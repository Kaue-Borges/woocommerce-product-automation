"""
Dashboard de Cadastro de Produtos - CHOICE
=============================================

Painel local (roda no seu PC, não precisa de internet além de acessar o
site da Choice) pra rodar a automação de cadastro sem precisar do
terminal. Você importa a planilha (.csv ou .xlsx) com os SKUs, aperta
Iniciar, e acompanha o progresso de cada produto em tempo real na tela.

Por trás, é a mesma automação por Playwright de antes: acha a pasta do
produto na rede, sobe a imagem principal, a galeria de fotos e todos os
arquivos de download (datasheet, manuais, desenho 3D, certificados...) e
salva o produto.

Padrões de pasta reconhecidos:
  A) Arquivos soltos com sufixo fixo:
       SKU_variante.jpg, SKU_foto2.jpg, SKU_datasheet.zip, SKU_3d.zip ...
  B) Arquivos soltos com nome descritivo:
       230010000159_HS200-1.png
       230010000159_Datasheet - Receptor HS200.zip
       230010000159_Manual de Pareamento para ....zip
       230010000159_Manual Série K_HS-K200 ....zip
       230010000159_Certificado Anatel - K200 ....zip
  C) Subpasta Fotos/ + Datasheets.zip / Desenho 3D.zip / Manual.zip na raiz

Nome do campo de download no site:
  - Se o produto tem só UM arquivo de um tipo, o campo leva o nome do
    tipo: "Datasheet", "Manual", "Desenho 3D", "Certificado".
  - Se tem MAIS DE UM do mesmo tipo (ex.: 2 manuais), cada um leva o nome
    descritivo do arquivo (o texto depois do SKU), pra o cliente saber
    qual é qual.

-----------------------------------------------------------------------
COMO USAR (Windows)
-----------------------------------------------------------------------
1. Instalar dependências (uma vez só):
       py -m pip install flask playwright openpyxl
       py -m playwright install chromium

2. Ajustar SITES abaixo (pasta raiz de cada loja) se a letra do drive de
   rede for diferente no seu PC. (Opcional) criar credenciais.json para
   login automático - ver comentário em CREDENCIAIS_FILE.

3. Rodar:
       py cadastro_dashboard.py

4. Abrir no navegador: http://localhost:5056

5. Na primeira vez que iniciar um lote, uma janela do Chrome separada vai
   abrir. Faça login manualmente no wp-admin da Choice nessa janela - o
   painel espera você logar e continua sozinho. Da próxima vez, reaproveita
   o login salvo.

6. Importe a planilha (.csv ou .xlsx, com uma coluna de SKU/Código - ou
   sem cabeçalho nenhum, só os códigos), confira a prévia dos SKUs
   encontrados, e clique em Iniciar.
-----------------------------------------------------------------------
"""

import io
import re
import threading
import time
from pathlib import Path

from flask import Flask, jsonify, request, send_file
from playwright.sync_api import sync_playwright, Page, FrameLocator, Locator


# ======================= CONFIG (ajuste aqui) =======================

# Lojas onde a automação pode rodar. Cada uma tem o próprio perfil do
# Chrome (login salvo separado) e a própria pasta raiz padrão na rede.
SITES = {
    "choice": {
        "nome": "CHOICE",
        "base_url": "https://choicetech.com.br",
        "root_folder": r"Z:\Choice\Geral\E-Commerce\CHOICE\Automação",
        "profile_dir": r".\chrome-profile-choice",
    },
    "sag": {
        "nome": "SAG",
        "base_url": "https://sagloja.devsplcomunic.com.br",
        "root_folder": r"Z:\Choice\Geral\E-Commerce\SAG\Automação",  # <- AJUSTE se for outra pasta
        "profile_dir": r".\chrome-profile-sag",
    },
}
SITE_PADRAO = "choice"

# Login automático (opcional). Crie um arquivo "credenciais.json" na mesma
# pasta do script, assim:
#   {"sag": {"usuario": "seu_usuario", "senha": "sua_senha"}}
# Se o arquivo não existir (ou a loja não estiver nele), o painel pede pra
# você logar manualmente na janela do Chrome, como antes.
CREDENCIAIS_FILE = Path(__file__).with_name("credenciais.json")

UPLOAD_TIMEOUT = 90000  # ms
PORT = 5056

# Extensões que contam como imagem e como arquivo de download
IMG_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".svg", ".gif"}
DOWNLOAD_EXTS = {".zip", ".pdf", ".rar", ".7z", ".step", ".stp", ".dwg", ".dxf"}

# Tipos de download reconhecidos pelo nome do arquivo, na ordem em que
# aparecem no site. (nome do campo, regex que identifica o tipo)
CATEGORIAS_DOWNLOAD = [
    ("Datasheet", re.compile(r"datasheet", re.IGNORECASE)),
    ("Manual", re.compile(r"manua", re.IGNORECASE)),
    ("Desenho 3D", re.compile(r"desenho|(?<![a-z0-9])3d(?![a-z0-9])", re.IGNORECASE)),
    ("Certificado", re.compile(r"certificad", re.IGNORECASE)),
]

# ======================================================================


# =========================================================================
# Identificação dos arquivos da pasta do produto
# =========================================================================

def find_product_folder(root: Path, sku: str) -> Path | None:
    sku_lower = sku.strip().lower()
    if not root.exists():
        raise RuntimeError(f"Pasta raiz não existe ou não está acessível: {root}")
    for path in root.rglob("*"):
        if path.is_dir() and path.name.lower() == sku_lower:
            return path
    return None


def _zip_na_raiz(folder: Path, *nomes: str) -> Path | None:
    """Procura um .zip pronto na raiz da pasta do produto, pelo nome exato
    (sem diferenciar maiúsculas/minúsculas). Ex.: 'Datasheets.zip',
    'Desenho 3D.zip', 'Manual.zip'."""
    alvos = {f"{n}.zip".lower() for n in nomes}
    for f in folder.iterdir():
        if f.is_file() and f.name.lower() in alvos:
            return f
    return None


def _descricao_do_arquivo(f: Path, sku: str) -> str:
    """'230010000159_Datasheet - Receptor HS200.zip' -> 'Datasheet - Receptor HS200'"""
    nome = re.sub(rf"^{re.escape(sku)}[\s_\-]*", "", f.stem, flags=re.IGNORECASE)
    nome = nome.strip(" _-")
    return nome or f.stem


def _categoria_do_arquivo(descricao: str) -> str | None:
    for nome, regex in CATEGORIAS_DOWNLOAD:
        if regex.search(descricao):
            return nome
    return None


def _montar_downloads(arquivos: list[tuple[str | None, Path]], sku: str) -> list[dict]:
    """Recebe [(categoria_forcada_ou_None, path)] e devolve a lista final de
    downloads [{label, categoria, path}], já ordenada e com nomes únicos.

    - categoria única no produto -> label = nome da categoria ("Manual")
    - categoria repetida         -> label = descrição do arquivo
    - sem categoria reconhecida  -> label = descrição do arquivo
    """
    itens = []
    vistos = set()
    for categoria, path in arquivos:
        if path is None or path in vistos:
            continue
        vistos.add(path)
        descricao = _descricao_do_arquivo(path, sku)
        if categoria is None:
            categoria = _categoria_do_arquivo(descricao)
        itens.append({"categoria": categoria, "descricao": descricao, "path": path})

    contagem = {}
    for it in itens:
        if it["categoria"]:
            contagem[it["categoria"]] = contagem.get(it["categoria"], 0) + 1

    for it in itens:
        if it["categoria"] and contagem[it["categoria"]] == 1:
            it["label"] = it["categoria"]
        else:
            it["label"] = it["descricao"]

    ordem = {nome: i for i, (nome, _) in enumerate(CATEGORIAS_DOWNLOAD)}
    itens.sort(key=lambda it: (ordem.get(it["categoria"], len(ordem)), it["label"].lower()))

    # garante que nenhum label se repete (a checagem pós-salvar usa o label)
    usados = {}
    for it in itens:
        base = it["label"]
        n = usados.get(base, 0) + 1
        usados[base] = n
        if n > 1:
            it["label"] = f"{base} ({n})"

    return [{"label": it["label"], "categoria": it["categoria"], "path": it["path"]} for it in itens]


def _classify_files_padrao_choice(folder: Path) -> dict:
    """Arquivos soltos na pasta do produto. Cobre:
      ..._variante                      -> imagem principal (se existir)
      ..._fotoN / _foto_N / _N / " N" / -N
                                        -> imagem numerada (a de menor
                                           número vira principal se não
                                           houver _variante; o resto vai
                                           pra galeria)
      ..._datasheet / _3d.zip / _manual -> downloads com sufixo fixo
      Datasheets.zip / Desenho 3D.zip / Manual.zip na raiz
      SKU_<qualquer nome>.zip/.pdf/...  -> download com nome descritivo
                                           (ex.: certificados, 2 manuais...)

    Exemplos cobertos:
      030PFFAR2FC_Variante.jpg, 030PFFAR2FC_Foto2.jpg
      99437_variante.svg, 99437_foto1.jpg, 99437_foto2.jpg
      RMCWA + DMR + DTO 1.jpg, RMCWA + DMR + DTO 2.jpg
      230010000159_HS200-1.png, 230010000159_Manual Série K_HS-K200.zip
    """
    files = [f for f in folder.iterdir() if f.is_file()]
    sku = folder.name

    def find_one(suffix_regex: str, ext: str | None = None):
        for f in files:
            if re.search(suffix_regex, f.stem, re.IGNORECASE):
                if ext and f.suffix.lower() != ext:
                    continue
                return f
        return None

    variante = find_one(r"_variante?$")
    datasheet = find_one(r"_datasheets?$") or _zip_na_raiz(folder, "Datasheets", "Datasheet")
    desenho3d = find_one(r"(_modelo)?_3d$", ext=".zip") or _zip_na_raiz(folder, "Desenho 3D", "Desenhos 3D")
    manual = find_one(r"_manual$") or _zip_na_raiz(folder, "Manual", "Manuais")

    ja_usados = {f for f in (variante, datasheet, desenho3d, manual) if f is not None}

    # ---- imagens numeradas ----
    numeradas = []
    for f in files:
        if f in ja_usados or f.suffix.lower() not in IMG_EXTS:
            continue
        m = re.search(r"(?:_foto_?|[\s_\-])(\d+)$", f.stem, re.IGNORECASE)
        if m:
            numeradas.append((int(m.group(1)), f))
    numeradas.sort(key=lambda t: t[0])

    fotos = [f for _, f in numeradas]
    if not variante and numeradas:
        variante = numeradas[0][1]
        fotos = [f for _, f in numeradas[1:]]

    # ---- downloads ----
    arquivos = [
        ("Datasheet", datasheet),
        ("Desenho 3D", desenho3d),
        ("Manual", manual),
    ]
    for f in sorted(files, key=lambda x: x.name.lower()):
        if f in ja_usados or f.suffix.lower() not in DOWNLOAD_EXTS:
            continue
        arquivos.append((None, f))

    return {
        "variante": variante,
        "fotos": fotos,
        "downloads": _montar_downloads(arquivos, sku),
    }


def _classify_files_padrao_subpastas(folder: Path) -> dict:
    """Fornecedor com subpasta Fotos/ e os .zip prontos na raiz da pasta do
    produto (Datasheets.zip / Desenho 3D.zip / Manual.zip). Qualquer outro
    .zip/.pdf solto na raiz também entra como download."""
    fotos_dir = folder / "Fotos"

    variante = None
    fotos = []
    if fotos_dir.is_dir():
        imagens = sorted(
            f for f in fotos_dir.iterdir()
            if f.is_file() and f.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp")
        )
        for f in imagens:
            if re.search(r"front$", f.stem, re.IGNORECASE):
                fotos.append(f)
            else:
                variante = f

    def zip_na_raiz(nome: str) -> Path | None:
        candidato = folder / f"{nome}.zip"
        return candidato if candidato.is_file() else None

    datasheet = zip_na_raiz("Datasheets")
    desenho3d = zip_na_raiz("Desenho 3D")
    manual = zip_na_raiz("Manual")

    arquivos = [
        ("Datasheet", datasheet),
        ("Desenho 3D", desenho3d),
        ("Manual", manual),
    ]
    usados = {datasheet, desenho3d, manual}
    for f in sorted(folder.iterdir(), key=lambda x: x.name.lower()):
        if f.is_file() and f not in usados and f.suffix.lower() in DOWNLOAD_EXTS:
            arquivos.append((None, f))

    return {
        "variante": variante,
        "fotos": fotos,
        "downloads": _montar_downloads(arquivos, folder.name),
    }


def classify_files(folder: Path) -> dict:
    """Detecta automaticamente qual padrão de pasta é esse e usa a lógica
    certa pra cada um."""
    if (folder / "Fotos").is_dir():
        return _classify_files_padrao_subpastas(folder)
    return _classify_files_padrao_choice(folder)


# =========================================================================
# Automação no WooCommerce
# =========================================================================

# A Choice usa o tema de admin "Adminify", que coloca o wp-admin dentro de
# um iframe. Outras lojas (ex.: SAG) podem usar o wp-admin normal. Isso é
# detectado automaticamente no começo de cada lote.
USAR_ADMINIFY = True


def admin_frame(page: Page) -> FrameLocator:
    if USAR_ADMINIFY:
        return page.locator("#frame-adminify-app--iframe").content_frame
    return page  # wp-admin normal: Page tem a mesma API de locator/get_by_role


def detectar_adminify(page: Page, log):
    global USAR_ADMINIFY
    try:
        page.locator("#frame-adminify-app--iframe").wait_for(state="attached", timeout=8000)
        USAR_ADMINIFY = True
        log("Admin com Adminify (iframe) detectado.")
    except Exception:
        USAR_ADMINIFY = False
        log("Admin padrão do WordPress detectado (sem iframe).")


def carregar_credenciais(site_key: str) -> dict | None:
    try:
        import json
        dados = json.loads(CREDENCIAIS_FILE.read_text(encoding="utf-8"))
        cred = dados.get(site_key)
        if cred and cred.get("usuario") and cred.get("senha"):
            return cred
    except Exception:
        pass
    return None


def esta_na_tela_de_login(page: Page) -> bool:
    try:
        return "wp-login.php" in page.url or page.locator("#loginform").count() > 0
    except Exception:
        return False


def tentar_login_automatico(page: Page, site_key: str, log) -> bool:
    cred = carregar_credenciais(site_key)
    if not cred:
        return False
    log("Fazendo login automático com credenciais.json...")
    try:
        page.locator("#user_login").fill(cred["usuario"], timeout=10000)
        page.locator("#user_pass").fill(cred["senha"], timeout=10000)
        page.locator("#wp-submit").click(timeout=10000)
        try:
            page.wait_for_load_state("networkidle", timeout=20000)
        except Exception:
            pass
    except Exception as e:
        log(f"  ! Login automático falhou: {e}")
        return False
    if esta_na_tela_de_login(page):
        log("  ! Login automático não passou (usuário/senha errados?).")
        return False
    log("Login automático OK.")
    return True


def clear_stale_selections(frame: FrameLocator, max_clicks: int = 25):
    for _ in range(max_clicks):
        btn = frame.get_by_role("button", name="Desmarcar").first
        try:
            if btn.count() == 0:
                return
            btn.click(timeout=2000)
        except Exception:
            return


def wait_enabled_and_click(locator: Locator, timeout_ms: int = UPLOAD_TIMEOUT):
    start = time.time()
    while (time.time() - start) * 1000 < timeout_ms:
        try:
            ready = locator.evaluate(
                """el => {
                    if (el.disabled) return false;
                    if (el.getAttribute('aria-disabled') === 'true') return false;
                    const s = window.getComputedStyle(el);
                    if (s.pointerEvents === 'none') return false;
                    return true;
                }"""
            )
        except Exception:
            ready = False
        if ready:
            break
        time.sleep(0.3)
    locator.click(timeout=15000)


def upload_and_confirm(page: Page, frame: FrameLocator, files: list[Path], confirm_button_name: str):
    clear_stale_selections(frame)
    try:
        frame.get_by_role("tab", name="Enviar arquivos").click(timeout=5000)
    except Exception:
        pass

    browse = frame.get_by_role("button", name="Selecionar arquivos")
    with page.expect_file_chooser(timeout=15000) as fc_info:
        browse.click()
    fc_info.value.set_files([str(f) for f in files])

    confirm = frame.get_by_role("button", name=confirm_button_name)
    confirm.wait_for(state="visible", timeout=UPLOAD_TIMEOUT)
    page.wait_for_timeout(800)
    wait_enabled_and_click(confirm, timeout_ms=UPLOAD_TIMEOUT)


def open_product_by_sku(page: Page, sku: str, base_url: str, log) -> bool:
    page.goto(f"{base_url}/wp-admin/edit.php?post_type=product")
    frame = admin_frame(page)

    search_box = frame.get_by_role("searchbox", name="Pesquisar produtos:")
    search_box.wait_for(state="visible", timeout=20000)
    search_box.click()
    search_box.fill(sku)
    frame.get_by_role("button", name="Pesquisar produtos").click()

    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass

    sku_pattern = re.compile(rf"^\s*{re.escape(sku)}\s*$")
    row = None
    for _ in range(24):
        sku_cell = frame.locator("td.column-sku").filter(has_text=sku_pattern)
        try:
            if sku_cell.count() > 0:
                row = sku_cell.first.locator("xpath=ancestor::tr")
                break
        except Exception:
            pass
        page.wait_for_timeout(500)

    if row is None:
        log(f"  !! Não encontrei nenhuma linha com o SKU exatamente igual a '{sku}'.")
        return False

    title_link = row.locator("a.row-title")
    if title_link.count() == 0:
        title_link = row.locator("strong a")
    if title_link.count() == 0:
        title_link = row.locator("a").first

    nome_produto = title_link.inner_text().strip()
    href = title_link.get_attribute("href")
    log(f"  Abrindo produto: {nome_produto}")

    if href:
        page.goto(href)
    else:
        title_link.click()

    frame = admin_frame(page)
    try:
        frame.locator("#title").wait_for(state="visible", timeout=20000)
    except Exception:
        log(f"  !! Não consegui confirmar que abriu a página de edição. URL atual: {page.url}")
        return False

    return True


def featured_image_is_set(frame: FrameLocator) -> bool:
    """Confere se a imagem principal realmente está definida (mostra uma
    miniatura de verdade, não só o texto/link de 'definir')."""
    try:
        return frame.locator("#set-post-thumbnail img").count() > 0
    except Exception:
        return False


def gallery_image_count(frame: FrameLocator) -> int:
    """Conta quantas imagens estão na galeria pelo campo escondido que o
    WooCommerce usa por baixo dos panos. Retorna -1 se não conseguir
    checar (não trava o processo, só não confirma)."""
    try:
        val = frame.locator("#product_image_gallery").input_value(timeout=3000)
        return len([x for x in val.split(",") if x.strip()])
    except Exception:
        return -1


def set_featured_image(page: Page, frame: FrameLocator, path: Path | None, log):
    if not path:
        log("  ! Aviso: imagem principal não encontrada - pulando.")
        return

    for tentativa in (1, 2):
        log(f"  Definindo imagem principal: {path.name} (tentativa {tentativa})")
        frame.locator("#set-post-thumbnail").click(timeout=15000)
        upload_and_confirm(page, frame, [path], confirm_button_name="Definir imagem do produto")
        page.wait_for_timeout(600)
        if featured_image_is_set(frame):
            log("  Imagem principal confirmada na tela.")
            return
        log("  ! A imagem principal não apareceu confirmada - tentando de novo...")

    raise RuntimeError("Imagem principal não confirmou depois de 2 tentativas.")


def add_gallery_images(page: Page, frame: FrameLocator, fotos: list[Path], log):
    if not fotos:
        log("  Nenhuma foto extra para a galeria - pulando.")
        return

    esperado = len(fotos)
    for tentativa in (1, 2):
        log(f"  Adicionando {esperado} foto(s) na galeria (tentativa {tentativa}).")
        frame.get_by_role("link", name="Adicionar galeria de imagens").click()
        upload_and_confirm(page, frame, fotos, confirm_button_name="Adicionar na galeria")
        page.wait_for_timeout(600)

        atual = gallery_image_count(frame)
        if atual == -1:
            log("  ! Não consegui confirmar a quantidade de fotos na galeria (seguindo mesmo assim).")
            return
        if atual >= esperado:
            log(f"  Galeria confirmada com {atual} foto(s) na tela.")
            return
        log(f"  ! Galeria mostrou {atual} foto(s), esperava pelo menos {esperado} - tentando de novo...")

    raise RuntimeError(f"Galeria não atingiu o número esperado de fotos ({esperado}) depois de 2 tentativas.")


def add_download_file(page: Page, frame: FrameLocator, nome: str, path: Path, log):
    log(f"  Adicionando arquivo de download '{nome}': {path.name}")
    label_inputs = frame.locator('input[name^="wcda_arquivos"][name$="[label]"]')
    index = label_inputs.count()

    frame.get_by_role("button", name="+ Adicionar Arquivo").click()
    label_input = label_inputs.nth(index)
    label_input.fill(nome)

    for tentativa in (1, 2):
        log(f"  Selecionando arquivo para '{nome}' (tentativa {tentativa})...")
        frame.get_by_role("button", name="Selecionar").nth(index).click()
        upload_and_confirm(page, frame, [path], confirm_button_name="Usar este arquivo")

        if _arquivo_download_confirmado(frame, index):
            log(f"  Arquivo '{nome}' confirmado (campo preenchido).")
            return
        log(f"  ! Não vi o campo do arquivo '{nome}' preenchido - tentando de novo...")

    raise RuntimeError(f"Não consegui confirmar que o arquivo '{nome}' foi anexado, mesmo após tentar de novo.")


def _arquivo_download_confirmado(frame: FrameLocator, index: int, timeout_ms: int = 5000) -> bool:
    """Confirma se o arquivo da linha 'index' foi vinculado de verdade,
    checando pelo NOME EXATO do campo (wcda_arquivos[index][...]) se algum
    campo dessa linha, além do de nome/label, recebeu um valor."""
    start = time.time()
    while (time.time() - start) * 1000 < timeout_ms:
        try:
            inputs = frame.locator(f'input[name^="wcda_arquivos[{index}]"]')
            for i in range(inputs.count()):
                campo = inputs.nth(i)
                nome_campo = campo.get_attribute("name") or ""
                if nome_campo.endswith("[label]"):
                    continue
                try:
                    valor = campo.input_value(timeout=1000)
                except Exception:
                    continue
                if valor and valor.strip():
                    return True
        except Exception:
            pass
        time.sleep(0.4)
    return False


def rolar_ate_botao_atualizar(page: Page, frame: FrameLocator, log) -> Locator:
    """Sobe a tela até o botão 'Atualizar' ficar visível de verdade antes de
    clicar."""
    botao = frame.locator("#publish")
    try:
        if botao.count() == 0:
            botao = frame.get_by_role("button", name="Atualizar")
    except Exception:
        botao = frame.get_by_role("button", name="Atualizar")
    botao = botao.first

    try:
        frame.locator("body").evaluate("() => window.scrollTo(0, 0)")
    except Exception:
        pass
    try:
        page.evaluate("() => window.scrollTo(0, 0)")
    except Exception:
        pass
    page.wait_for_timeout(400)

    for tentativa in range(1, 6):
        try:
            botao.scroll_into_view_if_needed(timeout=5000)
            botao.wait_for(state="visible", timeout=5000)
            na_tela = botao.evaluate(
                """el => {
                    const r = el.getBoundingClientRect();
                    const alt = window.innerHeight || document.documentElement.clientHeight;
                    return r.width > 0 && r.height > 0 && r.top >= 0 && r.bottom <= alt;
                }"""
            )
            if na_tela:
                log("  Botão 'Atualizar' visível na tela.")
                return botao
        except Exception:
            pass
        log(f"  Botão 'Atualizar' ainda fora da tela - rolando de novo (tentativa {tentativa})...")
        page.wait_for_timeout(600)

    log("  ! Não confirmei o botão 'Atualizar' na tela - tentando clicar mesmo assim.")
    return botao


def save_product(
    page: Page,
    frame: FrameLocator,
    log,
    esperado_imagem: bool,
    esperado_fotos: int,
    esperado_labels: list[str],
):
    log("  Salvando produto (Atualizar)...")
    botao = rolar_ate_botao_atualizar(page, frame, log)
    page.wait_for_timeout(300)
    botao.click(timeout=15000)

    # Espera mínima de segurança pro envio do formulário não ser cortado.
    page.wait_for_timeout(2500)

    try:
        handle = page.locator("#frame-adminify-app--iframe").element_handle(timeout=5000)
        real_frame = handle.content_frame() if handle else None
    except Exception:
        real_frame = None

    if real_frame is not None:
        try:
            real_frame.wait_for_load_state("networkidle", timeout=25000)
        except Exception:
            pass
    else:
        try:
            page.wait_for_load_state("networkidle", timeout=25000)
        except Exception:
            pass

    frame_apos = admin_frame(page)
    try:
        frame_apos.locator("#title").wait_for(state="visible", timeout=20000)
        log("  Página recarregou depois de salvar - confirmado.")
    except Exception:
        raise RuntimeError("A página não confirmou visualmente que recarregou depois de salvar.")

    # Checagem final: tudo persistiu depois de salvar?
    problemas = []

    if esperado_imagem and not featured_image_is_set(frame_apos):
        problemas.append("imagem principal não ficou definida")

    if esperado_fotos > 0:
        atual = gallery_image_count(frame_apos)
        if atual != -1 and atual < esperado_fotos:
            problemas.append(f"galeria ficou com {atual} foto(s), esperava {esperado_fotos}")

    if esperado_labels:
        label_inputs_apos = frame_apos.locator('input[name^="wcda_arquivos"][name$="[label]"]')
        valores_atuais = []
        for i in range(label_inputs_apos.count()):
            try:
                valores_atuais.append(label_inputs_apos.nth(i).input_value())
            except Exception:
                pass
        faltando = [l for l in esperado_labels if l not in valores_atuais]
        if faltando:
            problemas.append(f"campo(s) de download não persistiram: {', '.join(faltando)}")

    if problemas:
        raise RuntimeError("Depois de salvar: " + "; ".join(problemas))


def save_debug_screenshot(page: Page, sku: str) -> Path | None:
    try:
        debug_dir = Path("./debug_screenshots")
        debug_dir.mkdir(exist_ok=True)
        ts = time.strftime("%Y%m%d-%H%M%S")
        path = debug_dir / f"erro_{sku}_{ts}.png"
        page.screenshot(path=str(path), full_page=True)
        return path
    except Exception:
        return None


def process_sku(page: Page, sku: str, root: Path, base_url: str, log):
    log(f"--- Processando SKU {sku} ---")

    folder = find_product_folder(root, sku)
    if not folder:
        raise RuntimeError(f"Pasta não encontrada em {root} para o SKU {sku}.")
    log(f"  Pasta encontrada: {folder}")

    files = classify_files(folder)
    downloads = files["downloads"]

    log(f"  Imagem principal: {files['variante'].name if files['variante'] else 'NENHUMA'}")
    log(f"  Galeria: {[f.name for f in files['fotos']] if files['fotos'] else 'NENHUMA'}")
    if downloads:
        log(f"  Downloads ({len(downloads)}):")
        for d in downloads:
            log(f"    - '{d['label']}' <- {d['path'].name}")
    else:
        log("  Downloads: NENHUM")

    if not any([files["variante"], files["fotos"], downloads]):
        log(
            "  ! ATENÇÃO: não identifiquei NENHUM arquivo reconhecível nessa pasta. "
            "Confira se os nomes seguem um dos padrões esperados."
        )

    if not any(d["categoria"] == "Datasheet" for d in downloads):
        log("  ! Aviso: nenhum datasheet encontrado na pasta.")

    if not open_product_by_sku(page, sku, base_url, log):
        raise RuntimeError(f"Não consegui abrir o produto do SKU {sku} no WooCommerce.")

    frame = admin_frame(page)

    set_featured_image(page, frame, files["variante"], log)
    add_gallery_images(page, frame, files["fotos"], log)

    labels_adicionados = []
    for d in downloads:
        add_download_file(page, frame, d["label"], d["path"], log)
        labels_adicionados.append(d["label"])

    save_product(
        page, frame, log,
        esperado_imagem=bool(files["variante"]),
        esperado_fotos=len(files["fotos"]),
        esperado_labels=labels_adicionados,
    )
    log(f"  OK - SKU {sku} atualizado com sucesso.")


# =========================================================================
# Leitura da planilha (upload)
# =========================================================================

_COLUNAS_SKU_CONHECIDAS = ("sku", "código", "codigo", "cod", "cod.", "código cigam", "codigo cigam")
_SKU_REGEX = re.compile(r"^\d{5,}$")  # SKU real é só número, com pelo menos 5 dígitos


def _normalizar_valor_sku(val) -> str:
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def parse_skus_from_upload(filename: str, raw_bytes: bytes) -> list[str]:
    """Lê os SKUs de um arquivo enviado (.csv ou .xlsx/.xlsm). Detecta se
    tem cabeçalho ou não, e qual coluna usar."""
    ext = Path(filename).suffix.lower()
    linhas: list[list[str]] = []

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw_bytes), data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            linhas.append(["" if c is None else _normalizar_valor_sku(c) for c in row])

    elif ext == ".csv":
        import csv

        texto = None
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                texto = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if texto is None:
            raise ValueError("Não consegui ler a codificação do CSV.")

        amostra = texto[:2000]
        try:
            delim = csv.Sniffer().sniff(amostra, delimiters=",;\t").delimiter
        except Exception:
            delim = ","

        for row in csv.reader(io.StringIO(texto), delimiter=delim):
            linhas.append([c.strip() for c in row])

    else:
        raise ValueError(f"Formato não suportado: '{ext}' (use .csv ou .xlsx)")

    linhas = [l for l in linhas if any(c.strip() for c in l)]
    if not linhas:
        return []

    primeira = [c.strip().lower() for c in linhas[0]]
    tem_cabecalho = not all(_SKU_REGEX.match(c) for c in primeira if c)

    col_idx = 0
    if tem_cabecalho:
        col_idx = next((i for i, h in enumerate(primeira) if h in _COLUNAS_SKU_CONHECIDAS), 0)
        dados = linhas[1:]
    else:
        dados = linhas

    skus = []
    for row in dados:
        if col_idx < len(row) and row[col_idx]:
            skus.append(_normalizar_valor_sku(row[col_idx]))

    vistos = set()
    unicos = []
    for s in skus:
        if s not in vistos:
            vistos.add(s)
            unicos.append(s)
    return unicos


# =========================================================================
# Estado do lote / thread de automação
# =========================================================================

STATE_LOCK = threading.Lock()
STATE = {
    "running": False,
    "items": [],   # [{sku, status, detalhe}]
    "log": [],
    "need_login": False,
}


def log(msg: str):
    print(msg)
    with STATE_LOCK:
        STATE["log"].append(msg)
        if len(STATE["log"]) > 800:
            STATE["log"] = STATE["log"][-800:]


def update_item(sku: str, **kwargs):
    with STATE_LOCK:
        for item in STATE["items"]:
            if item["sku"] == sku:
                item.update(kwargs)
                return


def run_batch(skus: list[str], root_folder: str, site_key: str):
    site = SITES[site_key]
    base_url = site["base_url"]
    with STATE_LOCK:
        STATE["running"] = True
        STATE["items"] = [{"sku": s, "status": "Pendente", "detalhe": ""} for s in skus]
        STATE["log"] = []
        STATE["need_login"] = False

    root = Path(root_folder)

    try:
        with sync_playwright() as p:
            log(f"Loja: {site['nome']} ({base_url})")
            context = p.chromium.launch_persistent_context(
                user_data_dir=site["profile_dir"],
                headless=False,
                viewport={"width": 1440, "height": 900},
            )
            page = context.pages[0] if context.pages else context.new_page()

            log("Verificando login...")
            page.goto(f"{base_url}/wp-admin/edit.php?post_type=product")
            if esta_na_tela_de_login(page):
                tentar_login_automatico(page, site_key, log)
            if esta_na_tela_de_login(page):
                with STATE_LOCK:
                    STATE["need_login"] = True
                log(">>> Faça login manualmente na janela do Chrome que abriu. Aguardando...")
                while True:
                    time.sleep(1.5)
                    if not esta_na_tela_de_login(page):
                        break
                with STATE_LOCK:
                    STATE["need_login"] = False
                log("Login detectado, continuando...")
            page.goto(f"{base_url}/wp-admin/edit.php?post_type=product")
            detectar_adminify(page, log)

            for sku in skus:
                update_item(sku, status="Processando")
                try:
                    process_sku(page, sku, root, base_url, log)
                    update_item(sku, status="OK")
                except Exception as e:
                    shot = save_debug_screenshot(page, sku)
                    detalhe = str(e) + (f" (print: {shot})" if shot else "")
                    update_item(sku, status="Erro", detalhe=detalhe)
                    log(f"  !! ERRO no SKU {sku}: {e}")

            context.close()

    except Exception as e:
        log(f"ERRO GERAL: {e}")
    finally:
        with STATE_LOCK:
            STATE["running"] = False
        log("=== Lote finalizado ===")


# =========================================================================
# Flask
# =========================================================================

app = Flask(__name__)


@app.route("/")
def index():
    return HTML_PAGE


@app.route("/api/preview", methods=["POST"])
def api_preview():
    f = request.files.get("arquivo")
    if not f:
        return jsonify({"ok": False, "erro": "Nenhum arquivo enviado."})
    try:
        skus = parse_skus_from_upload(f.filename, f.read())
    except Exception as e:
        return jsonify({"ok": False, "erro": str(e)})
    return jsonify({"ok": True, "skus": skus})


@app.route("/api/iniciar", methods=["POST"])
def api_iniciar():
    with STATE_LOCK:
        if STATE["running"]:
            return jsonify({"ok": False, "erro": "Já tem um lote em andamento."})

    data = request.get_json()
    skus = data.get("skus") or []
    site_key = data.get("site") or SITE_PADRAO
    if site_key not in SITES:
        return jsonify({"ok": False, "erro": f"Loja desconhecida: {site_key}"})
    root_folder = (data.get("root_folder") or "").strip() or SITES[site_key]["root_folder"]
    if not skus:
        return jsonify({"ok": False, "erro": "Nenhum SKU informado."})

    thread = threading.Thread(target=run_batch, args=(skus, root_folder, site_key), daemon=True)
    thread.start()
    return jsonify({"ok": True, "total": len(skus)})


@app.route("/api/sites")
def api_sites():
    return jsonify({
        "padrao": SITE_PADRAO,
        "sites": [
            {"key": k, "nome": v["nome"], "base_url": v["base_url"], "root_folder": v["root_folder"]}
            for k, v in SITES.items()
        ],
    })


@app.route("/api/status")
def api_status():
    with STATE_LOCK:
        return jsonify(
            {
                "running": STATE["running"],
                "items": STATE["items"],
                "log": STATE["log"][-200:],
                "need_login": STATE["need_login"],
            }
        )


@app.route("/api/relatorio")
def api_relatorio():
    with STATE_LOCK:
        items = list(STATE["items"])
    ts = time.strftime("%Y%m%d-%H%M%S")
    path = Path(f"./resultado_lote_{ts}.txt")
    with open(path, "w", encoding="utf-8") as f:
        for item in items:
            f.write(f"{item['sku']}\t{item['status']}\t{item.get('detalhe', '')}\n")
    return send_file(path, as_attachment=True)


# =========================================================================
# Frontend
# =========================================================================

HTML_PAGE = """<!DOCTYPE html>
<html lang="pt-br">
<head>
<meta charset="UTF-8">
<title>Cadastro de Produtos</title>
<style>
  :root {
    --bg: #14161a; --panel: #1c1f26; --border: #2a2e37; --text: #e6e8eb;
    --muted: #8b93a1; --accent: #4f7cff; --green: #34c77b; --red: #ef5a5a;
    --yellow: #e0b23e;
  }
  * { box-sizing: border-box; }
  body {
    background: var(--bg); color: var(--text); font-family: -apple-system, Segoe UI, Roboto, sans-serif;
    margin: 0; padding: 32px; max-width: 980px; margin-inline: auto;
  }
  h1 { font-size: 20px; font-weight: 600; margin-bottom: 4px; }
  .sub { color: var(--muted); font-size: 13px; margin-bottom: 28px; }
  .panel {
    background: var(--panel); border: 1px solid var(--border); border-radius: 10px;
    padding: 20px; margin-bottom: 20px;
  }
  .panel h2 { font-size: 14px; font-weight: 600; margin: 0 0 14px; color: var(--muted); text-transform: uppercase; letter-spacing: .04em; }
  .dropzone {
    border: 2px dashed var(--border); border-radius: 8px; padding: 28px; text-align: center;
    color: var(--muted); cursor: pointer; transition: border-color .15s;
  }
  .dropzone:hover, .dropzone.drag { border-color: var(--accent); color: var(--text); }
  input[type=file] { display: none; }
  select {
    width: 100%; background: #0f1115; border: 1px solid var(--border);
    border-radius: 6px; padding: 8px 10px; color: var(--text); font-size: 13px; margin-top: 6px;
  }
  input[type=text] {
    width: 100%; background: #0f1115; border: 1px solid var(--border);
    border-radius: 6px; padding: 8px 10px; color: var(--text); font-size: 13px; margin-top: 6px;
  }
  label.field { display: block; font-size: 12px; color: var(--muted); margin-top: 14px; }
  button {
    background: var(--accent); color: white; border: none; border-radius: 6px;
    padding: 10px 18px; font-size: 14px; font-weight: 600; cursor: pointer; margin-top: 16px;
  }
  button:disabled { background: #3a3f4a; color: var(--muted); cursor: not-allowed; }
  button.secondary { background: #2a2e37; }
  #preview-list { margin-top: 14px; font-size: 13px; color: var(--muted); max-height: 140px; overflow-y: auto; }
  #preview-list span { display: inline-block; background: #262a33; border-radius: 4px; padding: 2px 8px; margin: 2px; color: var(--text); }
  table { width: 100%; border-collapse: collapse; font-size: 13px; margin-top: 4px; }
  th { text-align: left; color: var(--muted); font-weight: 500; padding: 6px 8px; border-bottom: 1px solid var(--border); }
  td { padding: 6px 8px; border-bottom: 1px solid var(--border); }
  .badge { padding: 2px 8px; border-radius: 999px; font-size: 11px; font-weight: 600; }
  .badge.Pendente { background: #2a2e37; color: var(--muted); }
  .badge.Processando { background: #2a3a5a; color: #7ea3ff; }
  .badge.OK { background: #1e3d2c; color: var(--green); }
  .badge.Erro { background: #3d1e1e; color: var(--red); }
  #log {
    background: #0b0d10; border: 1px solid var(--border); border-radius: 8px; padding: 12px;
    font-family: "SF Mono", Consolas, monospace; font-size: 12px; color: #a8b0bd;
    height: 220px; overflow-y: auto; white-space: pre-wrap;
  }
  .login-alert {
    background: #3d2e0f; border: 1px solid var(--yellow); color: var(--yellow);
    border-radius: 8px; padding: 10px 14px; font-size: 13px; margin-bottom: 14px;
  }
  .row { display: flex; gap: 10px; align-items: center; }
  .hidden { display: none !important; }
</style>
</head>
<body>

<h1>Cadastro de Produtos</h1>
<div class="sub">Importa a planilha de SKUs, revisa a prévia e inicia o lote.</div>

<div class="panel" id="panel-upload">
  <h2>1. Loja e planilha de SKUs</h2>
  <label class="field" style="margin-top:0;">Loja</label>
  <select id="site"></select>
  <div style="height:14px;"></div>
  <div class="dropzone" id="dropzone">
    Clique aqui ou arraste o arquivo (.csv ou .xlsx)
    <input type="file" id="file-input" accept=".csv,.xlsx,.xlsm">
  </div>
  <div id="preview-info" class="sub" style="margin-top:10px;"></div>
  <div id="preview-list"></div>

  <label class="field">Pasta raiz dos produtos na rede (deixe em branco pra usar o padrão)</label>
  <input type="text" id="root-folder" placeholder="">

  <button id="btn-iniciar" disabled>Iniciar cadastro</button>
</div>

<div class="panel hidden" id="panel-status">
  <h2>2. Progresso</h2>
  <div class="login-alert hidden" id="login-alert">
    Faça login manualmente na janela do Chrome que abriu. O painel continua sozinho depois que você logar.
  </div>
  <table>
    <thead><tr><th>SKU</th><th>Status</th><th>Detalhe</th></tr></thead>
    <tbody id="tbody-items"></tbody>
  </table>
  <div class="row" style="margin-top:14px;">
    <button class="secondary hidden" id="btn-relatorio">Baixar relatório</button>
  </div>
</div>

<div class="panel">
  <h2>Log</h2>
  <div id="log"></div>
</div>

<script>
let skusAtuais = [];
let sitesInfo = {};

const siteSelect = document.getElementById('site');
const rootInput = document.getElementById('root-folder');

(async () => {
  const res = await fetch('/api/sites');
  const data = await res.json();
  siteSelect.innerHTML = data.sites.map(s => {
    sitesInfo[s.key] = s;
    return '<option value="' + s.key + '">' + s.nome + ' — ' + s.base_url + '</option>';
  }).join('');
  siteSelect.value = data.padrao;
  atualizarPlaceholder();
})();

function atualizarPlaceholder() {
  const s = sitesInfo[siteSelect.value];
  rootInput.placeholder = s ? s.root_folder : '';
}
siteSelect.addEventListener('change', atualizarPlaceholder);

const dropzone = document.getElementById('dropzone');
const fileInput = document.getElementById('file-input');
const previewInfo = document.getElementById('preview-info');
const previewList = document.getElementById('preview-list');
const btnIniciar = document.getElementById('btn-iniciar');
const panelStatus = document.getElementById('panel-status');
const tbodyItems = document.getElementById('tbody-items');
const logDiv = document.getElementById('log');
const loginAlert = document.getElementById('login-alert');
const btnRelatorio = document.getElementById('btn-relatorio');

dropzone.addEventListener('click', () => fileInput.click());
dropzone.addEventListener('dragover', e => { e.preventDefault(); dropzone.classList.add('drag'); });
dropzone.addEventListener('dragleave', () => dropzone.classList.remove('drag'));
dropzone.addEventListener('drop', e => {
  e.preventDefault();
  dropzone.classList.remove('drag');
  if (e.dataTransfer.files.length) {
    fileInput.files = e.dataTransfer.files;
    enviarPreview(e.dataTransfer.files[0]);
  }
});
fileInput.addEventListener('change', () => {
  if (fileInput.files.length) enviarPreview(fileInput.files[0]);
});

async function enviarPreview(file) {
  previewInfo.textContent = 'Lendo ' + file.name + '...';
  const fd = new FormData();
  fd.append('arquivo', file);
  const res = await fetch('/api/preview', { method: 'POST', body: fd });
  const data = await res.json();
  if (!data.ok) {
    previewInfo.textContent = 'Erro: ' + data.erro;
    previewList.innerHTML = '';
    btnIniciar.disabled = true;
    return;
  }
  skusAtuais = data.skus;
  previewInfo.textContent = skusAtuais.length + ' SKU(s) encontrados em ' + file.name + ':';
  previewList.innerHTML = skusAtuais.map(s => '<span>' + s + '</span>').join('');
  btnIniciar.disabled = skusAtuais.length === 0;
}

btnIniciar.addEventListener('click', async () => {
  btnIniciar.disabled = true;
  const rootFolder = document.getElementById('root-folder').value;
  const res = await fetch('/api/iniciar', {
    method: 'POST',
    headers: { 'Content-Type': 'application/json' },
    body: JSON.stringify({ skus: skusAtuais, root_folder: rootFolder, site: siteSelect.value })
  });
  const data = await res.json();
  if (!data.ok) {
    alert(data.erro);
    btnIniciar.disabled = false;
    return;
  }
  panelStatus.classList.remove('hidden');
  btnRelatorio.classList.add('hidden');
  poll();
});

function renderItems(items) {
  tbodyItems.innerHTML = items.map(it => `
    <tr>
      <td>${it.sku}</td>
      <td><span class="badge ${it.status}">${it.status}</span></td>
      <td>${it.detalhe || ''}</td>
    </tr>
  `).join('');
}

async function poll() {
  const res = await fetch('/api/status');
  const data = await res.json();
  renderItems(data.items);
  logDiv.textContent = data.log.join('\\n');
  logDiv.scrollTop = logDiv.scrollHeight;
  loginAlert.classList.toggle('hidden', !data.need_login);

  if (data.running) {
    setTimeout(poll, 1200);
  } else {
    btnIniciar.disabled = false;
    if (data.items.length) btnRelatorio.classList.remove('hidden');
  }
}

btnRelatorio.addEventListener('click', () => {
  window.location.href = '/api/relatorio';
});

// se já tiver um lote rodando (ex: página recarregada), retoma o polling
(async () => {
  const res = await fetch('/api/status');
  const data = await res.json();
  if (data.items && data.items.length) {
    panelStatus.classList.remove('hidden');
    renderItems(data.items);
    logDiv.textContent = data.log.join('\\n');
    if (data.running) poll();
    else if (data.items.length) btnRelatorio.classList.remove('hidden');
  }
})();
</script>

</body>
</html>
"""


if __name__ == "__main__":
    print(f"\nAbra no navegador: http://localhost:{PORT}\n")
    app.run(host="127.0.0.1", port=PORT, debug=False)
